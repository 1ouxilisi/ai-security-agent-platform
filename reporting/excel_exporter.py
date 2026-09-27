#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 导出器 - 基于 openpyxl 生成多 Sheet 安全评估工作簿

Sheet:
- 执行摘要
- 漏洞清单（严重程度背景色标记）
- 端口清单
- 统计汇总（严重程度饼图 + 类型 Top10 + 按目标统计）

降级: openpyxl 不可用时生成多个 CSV 文件
"""
import csv
import os
from collections import Counter
from datetime import datetime
from typing import Any, Dict, List, Optional

from reporting._common import (
    SEVERITY_HEX,
    SEVERITY_LABELS_ZH,
    SEVERITY_ORDER,
    normalize_assessment,
    severity_counts,
)

try:
    import openpyxl
    from openpyxl.chart import PieChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# 表头样式常量
_HEADER_FILL = PatternFill("solid", fgColor="1E40AF")
_HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="CCCCCC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class ExcelExporter:
    """Excel 报告导出器"""

    def export(self, assessment_data: Dict[str, Any],
               output_path: str,
               template_config: Optional[Dict[str, Any]] = None) -> str:
        """导出 Excel 报告"""
        if not OPENPYXL_AVAILABLE:
            return self._fallback_csv(assessment_data, output_path)

        data = normalize_assessment(assessment_data)
        os.makedirs(os.path.dirname(os.path.abspath(output_path)) or ".", exist_ok=True)

        wb = openpyxl.Workbook()
        # 默认 sheet 改为执行摘要
        ws_summary = wb.active
        ws_summary.title = "执行摘要"
        self._build_summary_sheet(ws_summary, data)

        ws_vuln = wb.create_sheet("漏洞清单")
        self._build_vuln_sheet(ws_vuln, data)

        ws_port = wb.create_sheet("端口清单")
        self._build_port_sheet(ws_port, data)

        ws_stats = wb.create_sheet("统计汇总")
        self._build_stats_sheet(ws_stats, data)

        wb.save(output_path)
        return os.path.abspath(output_path)

    # ---------------- 样式工具 ----------------

    @staticmethod
    def _style_header(ws, row: int, ncols: int):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER
            cell.border = _BORDER

    @staticmethod
    def _auto_width(ws, min_width: int = 10, max_width: int = 50):
        for col_cells in ws.columns:
            length = min_width
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                v = "" if cell.value is None else str(cell.value)
                # 中文按 2 宽度计
                w = sum(2 if ord(ch) > 127 else 1 for ch in v)
                length = max(length, min(w, max_width))
            ws.column_dimensions[col_letter].width = length + 2

    # ---------------- 各 Sheet ----------------

    def _build_summary_sheet(self, ws, data):
        rows = [
            ("评估目标", data.get("target") or "-"),
            ("评估类型", data.get("assessment_type") or "-"),
            ("评估编号", data.get("assessment_id") or "-"),
            ("开始时间", self._fmt(data.get("started_at"))),
            ("完成时间", self._fmt(data.get("completed_at"))),
            ("综合风险分", data.get("overall_risk_score") or 0),
            ("综合风险等级", data.get("overall_risk_level") or "-"),
            ("漏洞总数", len(data["findings"])),
        ]
        ws.append(["项目", "内容"])
        self._style_header(ws, 1, 2)
        for k, v in rows:
            ws.append([k, v])
        # 关键发现
        ws.append([])
        ws.append(["关键发现"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        for kf in (data.get("key_findings") or ["无"]):
            ws.append(["•", kf])
        self._auto_width(ws)

    def _build_vuln_sheet(self, ws, data):
        headers = ["CVE", "漏洞名称", "严重程度", "描述", "证据",
                   "修复建议", "状态", "发现时间"]
        ws.append(headers)
        self._style_header(ws, 1, len(headers))

        for f in data["findings"]:
            sev = f["severity"]
            row = [
                f.get("cve") or "",
                f.get("title") or "",
                SEVERITY_LABELS_ZH.get(sev, sev),
                f.get("description") or "",
                f.get("evidence") or "",
                f.get("recommendation") or "",
                f.get("status") or "未修复",
                self._fmt(f.get("detected_at")),
            ]
            ws.append(row)
            # 严重程度单元格背景色
            color = SEVERITY_HEX.get(sev, "#6B7280").lstrip("#")
            fill = PatternFill("solid", fgColor=color)
            cell = ws.cell(row=ws.max_row, column=3)
            cell.fill = fill
            cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
            cell.alignment = _CENTER
            # 全表边框
            for c in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=c).border = _BORDER
                ws.cell(row=ws.max_row, column=c).alignment = Alignment(
                    vertical="top", wrap_text=True)

        ws.freeze_panes = "A2"
        self._auto_width(ws)

    def _build_port_sheet(self, ws, data):
        headers = ["端口", "协议", "服务", "版本", "状态"]
        ws.append(headers)
        self._style_header(ws, 1, len(headers))
        for p in data["ports"]:
            ws.append([
                str(p.get("port")),
                p.get("protocol"),
                p.get("service") or "",
                p.get("version") or "",
                p.get("status"),
            ])
            for c in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=c).border = _BORDER
        self._auto_width(ws)

    def _build_stats_sheet(self, ws, data):
        counts = severity_counts(data["findings"])

        # 严重程度统计（饼图数据源）
        ws.append(["严重程度", "数量"])
        self._style_header(ws, 1, 2)
        for sev in SEVERITY_ORDER:
            ws.append([SEVERITY_LABELS_ZH[sev], counts.get(sev, 0)])
        pie_first = ws.max_row - len(SEVERITY_ORDER) + 1
        pie_last = ws.max_row

        # 类型 Top 10
        ws.append([])
        type_start = ws.max_row + 1
        ws.append(["漏洞类型 Top10", "数量"])
        type_counter = Counter(f.get("category") or "未分类" for f in data["findings"])
        for cat, cnt in type_counter.most_common(10):
            ws.append([cat, cnt])
        type_end = ws.max_row

        # 按目标统计
        ws.append([])
        tgt_start = ws.max_row + 1
        ws.append(["目标", "漏洞数"])
        tgt_counter = Counter(f.get("target") or data.get("target") or "-"
                              for f in data["findings"])
        for tgt, cnt in tgt_counter.most_common(20):
            ws.append([tgt, cnt])
        tgt_end = ws.max_row

        # 饼图
        pie = PieChart()
        labels = Reference(ws, min_col=1, min_row=pie_first, max_row=pie_last)
        pdata = Reference(ws, min_col=2, min_row=1, max_row=pie_last)
        pie.add_data(pdata, titles_from_data=True)
        pie.set_categories(labels)
        pie.title = "漏洞严重程度分布"
        pie.height = 8
        pie.width = 12
        ws.add_chart(pie, "D2")

        self._auto_width(ws)

    # ---------------- 工具 ----------------

    @staticmethod
    def _fmt(ts) -> str:
        if not ts:
            return "-"
        try:
            return datetime.fromtimestamp(float(ts)).strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            return str(ts)

    # ---------------- 降级方案 ----------------

    def _fallback_csv(self, data, output_path: str) -> str:
        """openpyxl 不可用时拆分多个 CSV"""
        base = os.path.splitext(output_path)[0]
        data = normalize_assessment(data)
        sheets = {
            "漏洞清单": [["CVE", "名称", "严重程度", "描述", "证据", "修复建议", "状态", "发现时间"]]
            + [[f.get("cve", ""), f.get("title", ""), SEVERITY_LABELS_ZH[f["severity"]],
                f.get("description", ""), f.get("evidence", ""),
                f.get("recommendation", ""), f.get("status", ""),
                self._fmt(f.get("detected_at"))] for f in data["findings"]],
            "端口清单": [["端口", "协议", "服务", "版本", "状态"]]
            + [[str(p.get("port")), p.get("protocol"), p.get("service", ""),
                p.get("version", ""), p.get("status")] for p in data["ports"]],
        }
        out_paths = []
        for name, rows in sheets.items():
            p = f"{base}_{name}.csv"
            with open(p, "w", encoding="utf-8-sig", newline="") as fp:
                csv.writer(fp).writerows(rows)
            out_paths.append(os.path.abspath(p))
        return out_paths[0] if out_paths else output_path
