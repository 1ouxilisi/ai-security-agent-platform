#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word 导出器 - 基于 python-docx 生成可编辑的 .docx 安全评估报告

功能:
- 标题层级（Heading 1/2/3）、段落、表格、列表
- 中文字体（微软雅黑）、字号、颜色、表格样式
- 章节: 封面信息 / 执行摘要 / 漏洞详情表 / 端口列表表 / 修复建议 / 附录
- 降级: python-docx 不可用时生成 HTML（.doc 可被 Word 打开）
"""
import os
from datetime import datetime
from typing import Any, Dict, List, Optional

from reporting._common import (
    SEVERITY_HEX,
    SEVERITY_LABELS_ZH,
    SEVERITY_ORDER,
    normalize_assessment,
    severity_counts,
)

try:
    import docx
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm
    DOCX_AVAILABLE = True
except ImportError:  # 降级
    DOCX_AVAILABLE = False


class WordExporter:
    """Word 报告导出器"""

    CHINESE_FONT = "微软雅黑"

    def __init__(self):
        self.font_name = self.CHINESE_FONT if DOCX_AVAILABLE else "Arial"

    # ---------------- 公共入口 ----------------

    def export(self, assessment_data: Dict[str, Any],
               output_path: str,
               template_config: Optional[Dict[str, Any]] = None) -> str:
        """导出 Word 报告"""
        if not DOCX_AVAILABLE:
            return self._fallback_html(assessment_data, output_path, template_config)

        data = normalize_assessment(assessment_data)
        tpl = template_config or {}
        company = tpl.get("company_name") or "AI Hacking Agent"
        report_title = tpl.get("report_title") or "安全评估报告"
        report_no = tpl.get("report_no") or data.get("assessment_id") or ""
        sections = tpl.get("sections") or [
            "executive_summary", "vulnerability_details",
            "port_list", "remediation", "appendix",
        ]

        os.makedirs(os.path.dirname(os.path.abspath(output_path)) or ".", exist_ok=True)

        doc = Document()
        self._set_default_font(doc)

        # 封面信息
        self._build_cover(doc, data, tpl, company, report_title, report_no)

        for sec in sections:
            if sec == "executive_summary":
                self._section_executive_summary(doc, data)
            elif sec == "vulnerability_details":
                self._section_vuln_details(doc, data)
            elif sec == "port_list":
                self._section_port_list(doc, data)
            elif sec == "remediation":
                self._section_remediation(doc, data)
            elif sec == "appendix":
                self._section_appendix(doc, data)

        doc.save(output_path)
        return os.path.abspath(output_path)

    # ---------------- 样式工具 ----------------

    def _set_default_font(self, doc):
        """设置文档默认中文字体"""
        style = doc.styles["Normal"]
        style.font.name = self.font_name
        style.font.size = Pt(10.5)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), self.font_name)
        rfonts.set(qn("w:ascii"), self.font_name)
        rfonts.set(qn("w:hAnsi"), self.font_name)

    def _add_heading(self, doc, text: str, level: int = 1):
        """添加中文标题"""
        h = doc.add_heading(level=level)
        run = h.add_run(text)
        run.font.name = self.font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), self.font_name)
        if level == 1:
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif level == 2:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        else:
            run.font.size = Pt(12)
        return h

    def _add_paragraph(self, doc, text: str, bold: bool = False, size: int = 10.5):
        p = doc.add_paragraph()
        run = p.add_run(str(text))
        run.font.name = self.font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), self.font_name)
        run.font.size = Pt(size)
        run.bold = bold
        return p

    def _style_cell(self, cell, text: str, bold: bool = False,
                    color_hex: Optional[str] = None,
                    bg_hex: Optional[str] = None, size: int = 9):
        """设置单元格文字与底色"""
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(str(text))
        run.font.name = self.font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), self.font_name)
        run.font.size = Pt(size)
        run.bold = bold
        if color_hex:
            run.font.color.rgb = RGBColor.from_string(color_hex.lstrip("#"))
        if bg_hex:
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = tc_pr.makeelement(qn("w:shd"), {
                qn("w:val"): "clear",
                qn("w:color"): "auto",
                qn("w:fill"): bg_hex.lstrip("#"),
            })
            tc_pr.append(shd)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # ---------------- 各章节 ----------------

    def _build_cover(self, doc, data, tpl, company, report_title, report_no):
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(report_title)
        run.font.name = self.font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), self.font_name)
        run.font.size = Pt(28)
        run.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        srun = sub.add_run(company)
        srun.font.name = self.font_name
        srun._element.rPr.rFonts.set(qn("w:eastAsia"), self.font_name)
        srun.font.size = Pt(14)
        srun.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        doc.add_paragraph()
        meta = [
            ("评估目标", data.get("target") or "-"),
            ("评估类型", data.get("assessment_type") or "-"),
            ("报告编号", report_no or "-"),
            ("报告日期", datetime.now().strftime("%Y-%m-%d")),
            ("综合风险", f"{data.get('overall_risk_level')}（{data.get('overall_risk_score')} 分）"),
        ]
        table = doc.add_table(rows=len(meta), cols=2)
        table.style = "Light Grid Accent 1"
        for i, (k, v) in enumerate(meta):
            self._style_cell(table.rows[i].cells[0], k, bold=True, bg_hex="F1F5F9")
            self._style_cell(table.rows[i].cells[1], v)
        doc.add_page_break()

    def _section_executive_summary(self, doc, data):
        self._add_heading(doc, "一、执行摘要", 1)
        self._add_paragraph(doc, data.get("executive_summary") or "未提供执行摘要。")

        counts = severity_counts(data["findings"])
        self._add_heading(doc, "漏洞统计", 2)
        table = doc.add_table(rows=len(SEVERITY_ORDER) + 2, cols=2)
        table.style = "Light Grid Accent 1"
        self._style_cell(table.rows[0].cells[0], "严重程度", bold=True, bg_hex="1E40AF", color_hex="FFFFFF")
        self._style_cell(table.rows[0].cells[1], "数量", bold=True, bg_hex="1E40AF", color_hex="FFFFFF")
        for i, sev in enumerate(SEVERITY_ORDER, start=1):
            self._style_cell(table.rows[i].cells[0], SEVERITY_LABELS_ZH[sev],
                             bold=True, bg_hex=SEVERITY_HEX[sev], color_hex="FFFFFF")
            self._style_cell(table.rows[i].cells[1], str(counts.get(sev, 0)))
        total_row = len(SEVERITY_ORDER) + 1
        self._style_cell(table.rows[total_row].cells[0], "合计", bold=True)
        self._style_cell(table.rows[total_row].cells[1], str(len(data["findings"])), bold=True)

        if data.get("key_findings"):
            self._add_heading(doc, "关键发现", 2)
            for kf in data["key_findings"]:
                doc.add_paragraph(kf, style="List Bullet")

    def _section_vuln_details(self, doc, data):
        self._add_heading(doc, "二、漏洞详情", 1)
        findings = data["findings"]
        if not findings:
            self._add_paragraph(doc, "未发现漏洞。")
            return

        for sev in SEVERITY_ORDER:
            group = [f for f in findings if f["severity"] == sev]
            if not group:
                continue
            self._add_heading(doc, f"{SEVERITY_LABELS_ZH[sev]}（{len(group)} 项）", 2)

            headers = ["CVE", "漏洞名称", "严重程度", "描述", "证据", "修复建议"]
            table = doc.add_table(rows=len(group) + 1, cols=len(headers))
            table.style = "Light Grid Accent 1"
            for j, h in enumerate(headers):
                self._style_cell(table.rows[0].cells[j], h, bold=True,
                                 bg_hex="334155", color_hex="FFFFFF")
            for i, f in enumerate(group, start=1):
                self._style_cell(table.rows[i].cells[0], f.get("cve") or "-")
                self._style_cell(table.rows[i].cells[1], f.get("title") or "-")
                self._style_cell(table.rows[i].cells[2], SEVERITY_LABELS_ZH[sev],
                                 bold=True, bg_hex=SEVERITY_HEX[sev], color_hex="FFFFFF")
                self._style_cell(table.rows[i].cells[3],
                                 (f.get("description") or "-")[:300])
                self._style_cell(table.rows[i].cells[4],
                                 (f.get("evidence") or "-")[:200])
                self._style_cell(table.rows[i].cells[5],
                                 (f.get("recommendation") or "-")[:200])

    def _section_port_list(self, doc, data):
        self._add_heading(doc, "三、端口清单", 1)
        ports = data["ports"]
        if not ports:
            self._add_paragraph(doc, "未提供端口信息。")
            return
        headers = ["端口", "协议", "服务", "版本", "状态"]
        table = doc.add_table(rows=len(ports) + 1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        for j, h in enumerate(headers):
            self._style_cell(table.rows[0].cells[j], h, bold=True,
                             bg_hex="334155", color_hex="FFFFFF")
        for i, p in enumerate(ports, start=1):
            self._style_cell(table.rows[i].cells[0], str(p.get("port")))
            self._style_cell(table.rows[i].cells[1], str(p.get("protocol")))
            self._style_cell(table.rows[i].cells[2], str(p.get("service") or "-"))
            self._style_cell(table.rows[i].cells[3], str(p.get("version") or "-"))
            self._style_cell(table.rows[i].cells[4], str(p.get("status")))

    def _section_remediation(self, doc, data):
        self._add_heading(doc, "四、修复建议", 1)
        recs = data.get("recommendations") or []
        if not recs:
            self._add_paragraph(doc, "暂无修复建议。")
            return
        for rec in recs:
            doc.add_paragraph(rec, style="List Number")

    def _section_appendix(self, doc, data):
        self._add_heading(doc, "五、附录", 1)
        tools = data.get("tools_used") or []
        self._add_paragraph(doc, f"使用工具（{len(tools)}）：" + ("、".join(tools) or "无"), size=9)
        checks = data.get("checks_run") or []
        self._add_paragraph(doc, f"检测项（{len(checks)}）：" + ("、".join(checks) or "无"), size=9)
        self._add_paragraph(doc,
                            f"报告生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                            size=9)

    # ---------------- 降级方案 ----------------

    def _fallback_html(self, data, output_path, template_config) -> str:
        """python-docx 不可用时生成 HTML 内容，改后缀为 .doc"""
        out = output_path if output_path.endswith(".doc") else os.path.splitext(output_path)[0] + ".doc"
        tpl = template_config or {}
        company = tpl.get("company_name") or "AI Hacking Agent"
        rows = "".join(
            f"<tr><td>{f.get('cve','')}</td><td>{f.get('title','')}</td>"
            f"<td>{SEVERITY_LABELS_ZH[f['severity']]}</td>"
            f"<td>{f.get('description','')}</td></tr>"
            for f in data.get("findings", [])
        )
        html = f"""<html><head><meta charset="utf-8"><title>{tpl.get('report_title','报告')}</title></head>
<body><h1>{tpl.get('report_title','安全评估报告')}</h1>
<p>{company}</p>
<h2>执行摘要</h2><p>{data.get('executive_summary','')}</p>
<h2>漏洞详情</h2><table border="1"><tr><th>CVE</th><th>名称</th><th>严重</th><th>描述</th></tr>{rows}</table>
</body></html>"""
        with open(out, "w", encoding="utf-8") as fp:
            fp.write(html)
        return os.path.abspath(out)
